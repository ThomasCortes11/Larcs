#!/usr/bin/env python3
"""Importa el Excel del cliente al catálogo existente de LARCS."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

try:
    from openpyxl import load_workbook
    from PIL import Image as PILImage
except ImportError as error:
    raise SystemExit(
        "Faltan dependencias. Ejecuta: python -m pip install -r scripts/requirements-import.txt"
    ) from error


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "src" / "components" / "Excel" / "archivo excel.xlsx"
PRODUCTS_OUTPUT = ROOT / "public" / "products"
PRODUCTS_STAGING = ROOT / "public" / ".products-import-staging"
BACKUP_ROOT = ROOT / "backups" / "catalog"
MANIFEST_OUTPUT = ROOT / "src" / "data" / "products.json"
MANIFEST_STAGING = ROOT / "src" / "data" / ".products-import-staging.json"
DEFAULT_REPORT = ROOT / "reports" / "products-import-report.json"

SHEET_CATEGORIES = {
    "tacones": ("tacones", "Tacones"),
    "sandalias": ("sandalias", "Sandalias"),
    "botines": ("botines", "Botines"),
    "botas": ("botas", "Botas"),
    "flats": ("flats", "Flats")
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
IMAGE_MAX_EDGE = 2000
WEBP_QUALITY = 84
LEGACY_PRODUCT_FOLDERS = ("BOTAS", "BOTINES", "MOCASINES", "SANDALIAS", "TACONES")


def clean_reference(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_text = normalized.encode("ascii", "ignore").decode("ascii").lower()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_text).strip("-")
    return slug or "producto"


def category_for_sheet(title: str) -> tuple[str, str] | None:
    normalized = re.sub(r"\s+", " ", title).strip().casefold()
    return SHEET_CATEGORIES.get(normalized)


def find_header_row(worksheet) -> int | None:
    for row_index in range(1, min(worksheet.max_row, 12) + 1):
        first = str(worksheet.cell(row_index, 1).value or "").strip().casefold()
        second = str(worksheet.cell(row_index, 2).value or "").strip().casefold()
        if "referencia larcs" in first and "referencia" in second:
            return row_index
    return None


def parse_price(value: object) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(round(value))

    text = re.sub(r"[^0-9,.-]", "", str(value).strip())
    if not text:
        return None

    if "," in text and "." in text:
        decimal_separator = "," if text.rfind(",") > text.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        text = text.replace(thousands_separator, "").replace(decimal_separator, ".")
    elif "," in text:
        suffix = text.rsplit(",", 1)[1]
        text = text.replace(",", "." if len(suffix) != 3 else "")
    elif "." in text:
        suffix = text.rsplit(".", 1)[1]
        text = text.replace(".", "" if len(suffix) == 3 else ".")

    try:
        return int(round(float(text)))
    except ValueError:
        return None


def metadata_value(technical: str, label: str) -> str:
    prefix = f"{label.casefold()}:"
    for line in technical.splitlines():
        stripped = line.strip()
        if stripped.casefold().startswith(prefix):
            return stripped.split(":", 1)[1].strip()
    return ""


def parse_sizes(technical: str) -> list[str]:
    value = metadata_value(technical, "Tallas")
    range_match = re.search(r"(?<!\d)(\d{2})\s*[-–]\s*(\d{2})(?!\d)", value)
    if range_match:
        start, end = (int(number) for number in range_match.groups())
        if start <= end:
            return [str(size) for size in range(start, end + 1)]
    return re.findall(r"\d{2}", value)


def parse_colors(technical: str) -> list[str]:
    value = metadata_value(technical, "Color")
    return [color.strip() for color in value.split(",") if color.strip()]


def image_extension(image) -> str:
    del image
    return ".webp"


def remove_legacy_product_photos() -> int:
    removed = 0
    image_root = ROOT / "src" / "Img"
    for folder in LEGACY_PRODUCT_FOLDERS:
        directory = image_root / folder
        if not directory.is_dir():
            continue
        for path in directory.iterdir():
            if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS:
                path.unlink()
                removed += 1
    return removed


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        type=Path,
        default=DEFAULT_SOURCE,
        help="Ruta del Excel (por defecto: src/components/Excel/archivo excel.xlsx)."
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=DEFAULT_REPORT,
        help="Ruta del reporte (por defecto: reports/products-import-report.json)."
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Valida y genera el reporte sin reemplazar el catálogo ni las fotos."
    )
    args = parser.parse_args()
    source = args.source.resolve()
    report_path = args.report.resolve()

    if not source.is_file():
        print(f"No se encontró el archivo Excel: {source}", file=sys.stderr)
        return 1

    workbook = load_workbook(source, read_only=False, data_only=True)
    report = {
        "source": str(source),
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sheets": [],
        "totals": {},
        "productsWithoutImages": [],
        "productsWithoutPrice": [],
        "productsWithIncompleteInformation": [],
        "duplicateLarcsReferences": [],
        "errors": [],
        "warnings": []
    }
    products = []
    image_exports: list[tuple[object, Path]] = []
    source_image_bytes = 0
    optimized_image_bytes = 0
    larcs_references: dict[tuple[str, str], list[int]] = defaultdict(list)
    total_rows = 0
    total_images = 0

    for worksheet in workbook.worksheets:
        category = category_for_sheet(worksheet.title)
        if category is None:
            report["errors"].append(f"Hoja de productos no reconocida: {worksheet.title!r}")
            continue

        category_key, category_label = category
        header_row = find_header_row(worksheet)
        if header_row is None:
            report["errors"].append(f"No se encontró el encabezado en la hoja: {worksheet.title!r}")
            continue

        images_by_row: dict[int, list[tuple[int, int, object]]] = defaultdict(list)
        for image_index, image in enumerate(worksheet._images):
            anchor = getattr(image, "anchor", None)
            origin = getattr(anchor, "_from", None)
            if origin is None:
                report["warnings"].append(
                    f"Imagen sin ancla de celda en {worksheet.title!r}; se omitió."
                )
                continue
            images_by_row[origin.row + 1].append((origin.col, image_index, image))

        sheet_product_count = 0
        sheet_image_count = 0
        data_rows = []
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            reference_value = worksheet.cell(row_index, 1).value
            web_reference_value = worksheet.cell(row_index, 2).value
            technical_value = worksheet.cell(row_index, 3).value
            price_value = worksheet.cell(row_index, 4).value
            if all(value is None for value in (reference_value, web_reference_value, technical_value, price_value)):
                continue

            total_rows += 1
            sheet_product_count += 1
            reference = clean_reference(reference_value)
            web_reference = clean_reference(web_reference_value)
            technical = "" if technical_value is None else str(technical_value)
            price = parse_price(price_value)
            colors = parse_colors(technical)
            sizes = parse_sizes(technical)
            missing_fields = []

            if not reference:
                missing_fields.append("Referencia Larcs")
            if not web_reference:
                missing_fields.append("Referencia Página Web")
            if not technical.strip():
                missing_fields.append("Fecha Técnica")
            if not colors:
                missing_fields.append("Color")
            if not sizes:
                missing_fields.append("Tallas")
            if price is None:
                missing_fields.append("Precio Venta")
                report["productsWithoutPrice"].append(
                    {"sheet": category_label, "row": row_index, "name": web_reference}
                )
            if missing_fields:
                report["productsWithIncompleteInformation"].append(
                    {
                        "sheet": category_label,
                        "row": row_index,
                        "name": web_reference,
                        "missingFields": missing_fields
                    }
                )

            if not web_reference:
                report["errors"].append(
                    f"Falta el nombre del producto en {worksheet.title!r}, fila {row_index}."
                )
                continue
            if price is None:
                report["errors"].append(
                    f"Precio ausente o inválido en {worksheet.title!r}, fila {row_index}."
                )
                continue

            product_slug = f"{category_key}-{slugify(web_reference)}-{row_index:03d}"
            product_images = sorted(images_by_row.get(row_index, []), key=lambda item: (item[0], item[1]))
            if not product_images:
                report["productsWithoutImages"].append(
                    {"sheet": category_label, "row": row_index, "name": web_reference}
                )
                missing_fields.append("Fotografías")

            if missing_fields:
                incomplete_entry = next(
                    (
                        entry
                        for entry in report["productsWithIncompleteInformation"]
                        if entry["sheet"] == category_label and entry["row"] == row_index
                    ),
                    None
                )
                if incomplete_entry is None:
                    report["productsWithIncompleteInformation"].append(
                        {
                            "sheet": category_label,
                            "row": row_index,
                            "name": web_reference,
                            "missingFields": missing_fields
                        }
                    )
                else:
                    incomplete_entry["missingFields"] = missing_fields

            image_urls = []
            for image_number, (_, _, image) in enumerate(product_images, start=1):
                extension = image_extension(image)
                relative_path = Path(category_key) / product_slug / f"{image_number}{extension}"
                image_urls.append(f"/products/{relative_path.as_posix()}")
                image_exports.append((image, PRODUCTS_STAGING / relative_path))
                sheet_image_count += 1

            data_rows.append({"row": row_index, "reference": reference})
            if reference:
                larcs_references[(category_label, reference.casefold())].append(row_index)

            products.append(
                {
                    "id": product_slug,
                    "slug": product_slug,
                    "reference": reference,
                    "webReference": web_reference,
                    "name": web_reference,
                    "category": category_key,
                    "categoryLabel": category_label,
                    "price": price,
                    "description": technical,
                    "features": [],
                    "popularity": 0,
                    "isNew": False,
                    "isPromo": False,
                    "imageUrls": image_urls,
                    "variant": {
                        "color": metadata_value(technical, "Color"),
                        "colors": colors,
                        "sizes": sizes,
                        "stock": None
                    }
                }
            )

        image_rows_without_product = sorted(set(images_by_row) - {row for row, _ in data_rows})
        for row_index in image_rows_without_product:
            report["warnings"].append(
                f"Hay imágenes ancladas a una fila sin producto en {worksheet.title!r}, fila {row_index}."
            )

        total_images += len(worksheet._images)
        report["sheets"].append(
            {
                "name": worksheet.title,
                "category": category_label,
                "productsFound": sheet_product_count,
                "imagesFound": len(worksheet._images),
                "imagesAssociated": sheet_image_count,
                "rowsWithImages": len(images_by_row),
                "imageRowsWithoutProducts": image_rows_without_product
            }
        )

    for (category_label, reference), rows in larcs_references.items():
        if len(rows) > 1:
            report["duplicateLarcsReferences"].append(
                {"category": category_label, "reference": reference, "rows": rows}
            )

    report["totals"] = {
        "productsFound": total_rows,
        "productsImported": len(products),
        "imagesFound": total_images,
        "imagesExported": len(image_exports),
        "productsWithoutImages": len(report["productsWithoutImages"]),
        "productsWithoutPrice": len(report["productsWithoutPrice"]),
        "productsWithIncompleteInformation": len(report["productsWithIncompleteInformation"]),
        "productsByCategory": {
            sheet["category"]: sheet["productsFound"] for sheet in report["sheets"]
        },
        "imagesByCategory": {
            sheet["category"]: sheet["imagesFound"] for sheet in report["sheets"]
        }
    }

    if report["errors"]:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        print("Importación cancelada; el catálogo actual no se modificó.", file=sys.stderr)
        return 1

    if args.dry_run:
        report["dryRun"] = True
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8"
        )
        print(json.dumps({"validacion": "correcta", "report": str(report_path), **report["totals"]}, ensure_ascii=False, indent=2))
        return 0

    if PRODUCTS_STAGING.exists():
        shutil.rmtree(PRODUCTS_STAGING)
    PRODUCTS_STAGING.mkdir(parents=True, exist_ok=True)
    for image, destination in image_exports:
        destination.parent.mkdir(parents=True, exist_ok=True)
        source_bytes = image._data()
        source_image_bytes += len(source_bytes)
        with PILImage.open(BytesIO(source_bytes)) as source_image:
            optimized_image = source_image.copy()
        if max(optimized_image.size) > IMAGE_MAX_EDGE:
            optimized_image.thumbnail(
                (IMAGE_MAX_EDGE, IMAGE_MAX_EDGE),
                PILImage.Resampling.LANCZOS
            )
        if optimized_image.mode not in ("RGB", "RGBA"):
            optimized_image = optimized_image.convert(
                "RGBA" if "A" in optimized_image.getbands() else "RGB"
            )
        optimized_image.save(
            destination,
            format="WEBP",
            quality=WEBP_QUALITY,
            method=6
        )
        optimized_image_bytes += destination.stat().st_size

    MANIFEST_STAGING.parent.mkdir(parents=True, exist_ok=True)
    MANIFEST_STAGING.write_text(
        json.dumps(products, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8"
    )

    backup_directory = None
    if PRODUCTS_OUTPUT.exists() or MANIFEST_OUTPUT.exists():
        backup_name = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
        backup_directory = BACKUP_ROOT / backup_name
        backup_directory.mkdir(parents=True, exist_ok=False)
        if MANIFEST_OUTPUT.exists():
            shutil.copy2(MANIFEST_OUTPUT, backup_directory / "products.json")
        if PRODUCTS_OUTPUT.exists():
            PRODUCTS_OUTPUT.replace(backup_directory / "products")

    try:
        PRODUCTS_STAGING.replace(PRODUCTS_OUTPUT)
        MANIFEST_STAGING.replace(MANIFEST_OUTPUT)
    except OSError:
        if PRODUCTS_OUTPUT.exists():
            shutil.rmtree(PRODUCTS_OUTPUT)
        previous_products = backup_directory / "products" if backup_directory else None
        if previous_products and previous_products.exists():
            previous_products.replace(PRODUCTS_OUTPUT)
        previous_manifest = backup_directory / "products.json" if backup_directory else None
        if previous_manifest and previous_manifest.exists():
            shutil.copy2(previous_manifest, MANIFEST_OUTPUT)
        raise

    report["legacyProductPhotosRemoved"] = remove_legacy_product_photos()
    report["backupPath"] = str(backup_directory.relative_to(ROOT)) if backup_directory else None
    report["imageOptimization"] = {
        "format": "WebP",
        "maximumEdgePixels": IMAGE_MAX_EDGE,
        "quality": WEBP_QUALITY,
        "sourceBytes": source_image_bytes,
        "optimizedBytes": optimized_image_bytes,
        "reductionPercent": round(
            (1 - optimized_image_bytes / source_image_bytes) * 100, 1
        )
        if source_image_bytes
        else 0
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8"
    )

    print(json.dumps({"report": str(report_path), **report["totals"], "legacyProductPhotosRemoved": report["legacyProductPhotosRemoved"]}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())